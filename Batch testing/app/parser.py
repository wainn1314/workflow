"""Excel / CSV 测试集解析。"""
import ast
import csv
import io
import json
from typing import Any, Dict, List

REQUIRED_COLUMNS = ["test_case_id", "conversation", "expected_output", "scenario", "difficulty"]


def parse_file(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    low = (filename or "").lower()
    if low.endswith((".xlsx", ".xlsm", ".xls")):
        return parse_excel(file_bytes)
    return parse_csv(file_bytes)


def parse_csv(data: bytes) -> List[Dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for i, raw in enumerate(reader, start=1):
        cleaned = {str(k).strip(): (v or "").strip() for k, v in raw.items() if k is not None}
        if not any(cleaned.values()):
            continue
        rows.append(normalize_case(cleaned, i))
    return rows


def parse_excel(data: bytes) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return []
        header = [str(h).strip() if h is not None else "" for h in header]
        rows = []
        for i, values in enumerate(it, start=1):
            cleaned = {
                (str(k).strip() if k is not None else ""): ("" if v is None else str(v).strip())
                for k, v in zip(header, values) if k
            }
            if not any(cleaned.values()):
                continue
            rows.append(normalize_case(cleaned, i))
        return rows
    finally:
        wb.close()


def normalize_case(row: Dict[str, str], idx: int) -> Dict[str, Any]:
    missing = [c for c in REQUIRED_COLUMNS if not row.get(c)]
    if missing:
        raise ValueError(f"第 {idx} 行缺少必填字段: {missing}")
    conversation = parse_conversation(row["conversation"])
    if not conversation:
        raise ValueError(f"第 {idx} 行 conversation 无法解析为 messages 数组")
    if not any(m.get("role") == "user" for m in conversation):
        raise ValueError(f"第 {idx} 行 conversation 中缺少 user 消息")
    return {
        "ref_id": str(row["test_case_id"]),
        "scenario": str(row["scenario"]),
        "difficulty": str(row["difficulty"]),
        "conversation": conversation,
        "expected_output": str(row["expected_output"]),
    }


def parse_conversation(raw: str) -> List[Dict[str, str]]:
    raw = (raw or "").strip()
    if not raw:
        return []
    data = None
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        try:
            data = ast.literal_eval(raw)
        except (ValueError, SyntaxError):
            return []
    if isinstance(data, dict) and isinstance(data.get("messages"), list):
        data = data["messages"]
    if not isinstance(data, list):
        return []
    out = []
    for m in data:
        if isinstance(m, dict) and m.get("role") and m.get("content") is not None:
            out.append({"role": str(m["role"]), "content": str(m["content"])})
        else:
            return []
    return out
